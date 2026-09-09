import asyncio
import json
import re
import tempfile
import threading
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook
from config.settings import SalesReport
from utils.extract import extract_statement
from utils.helpers import (
    export_combined_results,
    extract_zip,
    find_pdfs,
    reconcile_with_statement,
)
from utils.gemini_limits import GeminiQuotaCoordinator
from utils.sales_helpers import process_sales_zip, write_to_excel_with_categories


class FakeAsyncClient:
    def __init__(self, models=None):
        self.models = models

    async def aclose(self):
        pass


class FakeGenaiClient:
    def __init__(self, models=None):
        self.aio = FakeAsyncClient(models)

    def close(self):
        pass


class FileProcessingTests(unittest.TestCase):
    def test_extracts_and_finds_pdfs_at_any_depth(self):
        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            archive = root / "reports.zip"
            with zipfile.ZipFile(archive, "w") as zipped:
                zipped.writestr("root.PDF", b"pdf")
                zipped.writestr("a/b/nested.pdf", b"pdf")
                zipped.writestr("__MACOSX/._root.PDF", b"metadata")
                zipped.writestr("a/b/._nested.pdf", b"metadata")

            extracted = extract_zip(archive, root / "extracted")
            self.assertEqual(
                [path.name for path in find_pdfs(extracted)],
                ["nested.pdf", "root.PDF"],
            )

    def test_rejects_unsafe_or_oversized_zip(self):
        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            unsafe = root / "unsafe.zip"
            with zipfile.ZipFile(unsafe, "w") as zipped:
                zipped.writestr("../escape.pdf", b"pdf")
            with self.assertRaisesRegex(ValueError, "unsafe"):
                extract_zip(unsafe, root / "unsafe-output")

            large = root / "large.zip"
            with zipfile.ZipFile(large, "w") as zipped:
                zipped.writestr("large.pdf", b"1234")
            with self.assertRaisesRegex(ValueError, "safety limit"):
                extract_zip(large, root / "large-output", max_bytes=3)

    def test_empty_reconciliation_is_rejected(self):
        with tempfile.TemporaryDirectory() as workspace:
            output = Path(workspace) / "report.xlsx"
            with self.assertRaisesRegex(ValueError, "No invoices"):
                export_combined_results([], [], output)
            self.assertFalse(output.exists())

    def test_reconciliation_writes_to_the_requested_path(self):
        invoice = {"vendor_name": "Vendor", "reconciled": True}
        with tempfile.TemporaryDirectory() as workspace:
            output = Path(workspace) / "unique-report.xlsx"
            actual = export_combined_results([invoice], [], output)

            self.assertEqual(actual, output)
            self.assertEqual(load_workbook(output).sheetnames, ["Matched"])


class SalesProcessingTests(unittest.TestCase):
    def test_empty_and_all_failed_batches_are_rejected(self):
        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            with self.assertRaisesRegex(ValueError, "No PDF"):
                process_sales_zip(root, root / "output.xlsx")

            (root / "report.pdf").write_bytes(b"pdf")
            with self.assertLogs("utils.sales_helpers", level="INFO"):
                with (
                    patch(
                        "utils.sales_helpers.create_genai_client",
                        return_value=FakeGenaiClient(),
                    ),
                    patch("utils.sales_helpers.extract_sales_data", side_effect=RuntimeError),
                ):
                    with self.assertRaisesRegex(RuntimeError, "Could not process any"):
                        process_sales_zip(root, root / "output.xlsx")

    def test_partial_batch_reports_processed_and_failed_counts(self):
        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            (root / "one.pdf").write_bytes(b"pdf")
            (root / "two.pdf").write_bytes(b"pdf")
            output = root / "output.xlsx"
            progress = []

            with self.assertLogs("utils.sales_helpers", level="INFO") as logs:
                with (
                    patch(
                        "utils.sales_helpers.create_genai_client",
                        return_value=FakeGenaiClient(),
                    ),
                    patch(
                        "utils.sales_helpers.extract_sales_data",
                        side_effect=[{"Date": "2026-07-01"}, RuntimeError("bad PDF")],
                    ),
                    patch("utils.sales_helpers.write_to_excel_with_categories") as write,
                ):
                    processed, failed = process_sales_zip(
                        root,
                        output,
                        lambda current, total, name, status: progress.append(
                            (current, total, name, status)
                        ),
                    )

            self.assertEqual((processed, failed), (1, 1))
            self.assertEqual(write.call_args.args[1], output)
            self.assertEqual([event[0] for event in progress], [1, 2])
            self.assertEqual(
                {(event[2], event[3]) for event in progress},
                {("one.pdf", "processed"), ("two.pdf", "failed")},
            )
            self.assertTrue(any("processed 1/2: one.pdf" in line for line in logs.output))
            self.assertTrue(any("failed 2/2: two.pdf" in line for line in logs.output))

    def test_per_job_concurrency_and_progress(self):
        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            for index in range(10):
                (root / f"{index}.pdf").write_bytes(b"pdf")

            active = peak = 0
            progress = []

            async def extract(*args):
                nonlocal active, peak
                active += 1
                peak = max(peak, active)
                await asyncio.sleep(0.01)
                active -= 1
                return {"Date": "2026-07-01"}

            with (
                patch("utils.sales_helpers.create_genai_client", return_value=FakeGenaiClient()),
                patch("utils.sales_helpers.extract_sales_data", side_effect=extract),
                patch("utils.sales_helpers.write_to_excel_with_categories"),
            ):
                result = process_sales_zip(
                    root,
                    root / "output.xlsx",
                    lambda *event: progress.append(event),
                )

            self.assertEqual(result, (10, 0))
            self.assertGreater(peak, 1)
            self.assertLessEqual(peak, 6)
            self.assertEqual(len(progress), 10)
            self.assertEqual([event[0] for event in progress], list(range(1, 11)))

    def test_two_jobs_share_global_coordinator(self):
        class Models:
            def __init__(self):
                self.active = 0
                self.peak = 0
                self.lock = threading.Lock()

            async def generate_content(self, **kwargs):
                with self.lock:
                    self.active += 1
                    self.peak = max(self.peak, self.active)
                await asyncio.sleep(0.02)
                with self.lock:
                    self.active -= 1
                parsed = SimpleNamespace(model_dump=lambda: {"Date": "2026-07-01"})
                usage = SimpleNamespace(prompt_token_count=1_000)
                return SimpleNamespace(parsed=parsed, usage_metadata=usage)

        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            folders = [root / "first", root / "second"]
            for folder in folders:
                folder.mkdir()
                for index in range(5):
                    (folder / f"{index}.pdf").write_bytes(b"pdf")

            models = Models()
            coordinator = GeminiQuotaCoordinator(
                1_000, 10_000_000, headroom=1, max_concurrency=3
            )
            with (
                patch("utils.sales_helpers.get_gemini_coordinator", return_value=coordinator),
                patch(
                    "utils.sales_helpers.create_genai_client",
                    side_effect=lambda: FakeGenaiClient(models),
                ),
                patch("utils.sales_helpers.write_to_excel_with_categories"),
                self.assertLogs("utils.sales_helpers", level="INFO") as logs,
            ):
                with ThreadPoolExecutor(max_workers=2) as pool:
                    results = list(
                        pool.map(
                            lambda folder: process_sales_zip(
                                folder, folder / "output.xlsx"
                            ),
                            folders,
                        )
                    )

            job_ids = {
                match.group(1)
                for line in logs.output
                if (match := re.search(r"Sales job ([0-9a-f]{8}) queued", line))
            }
            self.assertEqual(results, [(5, 0), (5, 0)])
            self.assertGreater(models.peak, 1)
            self.assertLessEqual(models.peak, 3)
            self.assertEqual(len(job_ids), 2)

    def test_retries_once_in_progress_and_non_retryable_fails_immediately(self):
        from google.genai.errors import ClientError

        class Models:
            def __init__(self, error_code, failures=2):
                self.calls = 0
                self.error_code = error_code
                self.failures = failures

            async def generate_content(self, **kwargs):
                self.calls += 1
                if self.error_code and self.calls <= self.failures:
                    raise ClientError(self.error_code, {"message": "test"})
                parsed = SimpleNamespace(model_dump=lambda: {"Date": "2026-07-01"})
                usage = SimpleNamespace(prompt_token_count=1_000)
                return SimpleNamespace(parsed=parsed, usage_metadata=usage)

        with tempfile.TemporaryDirectory() as workspace:
            root = Path(workspace)
            (root / "report.pdf").write_bytes(b"pdf")
            progress = []
            coordinator = GeminiQuotaCoordinator(1_000, 10_000_000, headroom=1)

            retrying = Models(429)
            with (
                patch("utils.sales_helpers.get_gemini_coordinator", return_value=coordinator),
                patch(
                    "utils.sales_helpers.create_genai_client",
                    return_value=FakeGenaiClient(retrying),
                ),
                patch("utils.sales_helpers.retry_delay", return_value=0),
                patch("utils.sales_helpers.write_to_excel_with_categories"),
            ):
                self.assertEqual(
                    process_sales_zip(
                        root, root / "retry.xlsx", lambda *event: progress.append(event)
                    ),
                    (1, 0),
                )

            self.assertEqual(retrying.calls, 3)
            self.assertEqual(len(progress), 1)

            rejected = Models(400)
            with (
                patch("utils.sales_helpers.get_gemini_coordinator", return_value=coordinator),
                patch(
                    "utils.sales_helpers.create_genai_client",
                    return_value=FakeGenaiClient(rejected),
                ),
                patch("utils.sales_helpers.write_to_excel_with_categories"),
            ):
                with self.assertRaisesRegex(RuntimeError, "Could not process any"):
                    process_sales_zip(root, root / "rejected.xlsx")
            self.assertEqual(rejected.calls, 1)

            exhausted = Models(429, failures=6)
            progress.clear()
            with (
                patch("utils.sales_helpers.get_gemini_coordinator", return_value=coordinator),
                patch(
                    "utils.sales_helpers.create_genai_client",
                    return_value=FakeGenaiClient(exhausted),
                ),
                patch("utils.sales_helpers.retry_delay", return_value=0),
                patch("utils.sales_helpers.write_to_excel_with_categories"),
            ):
                with self.assertRaisesRegex(RuntimeError, "Could not process any"):
                    process_sales_zip(
                        root,
                        root / "exhausted.xlsx",
                        lambda *event: progress.append(event),
                    )
            self.assertEqual(exhausted.calls, 6)
            self.assertEqual(progress[0][3], "failed")

    def test_out_of_order_rows_are_sorted_in_workbook(self):
        def sales_row(day):
            return {
                field: (
                    date(2026, 7, day)
                    if field == "Date"
                    else None if field == "Notes" else 0
                )
                for field in SalesReport.model_fields
            }

        with tempfile.TemporaryDirectory() as workspace:
            output = Path(workspace) / "sorted.xlsx"
            write_to_excel_with_categories(
                pd.DataFrame([sales_row(2), sales_row(1)]), output
            )
            worksheet = load_workbook(output).active
            self.assertEqual(
                [worksheet["A3"].value.date(), worksheet["A4"].value.date()],
                [date(2026, 7, 1), date(2026, 7, 2)],
            )


class GeminiQuotaTests(unittest.TestCase):
    def test_headroom_is_applied(self):
        coordinator = GeminiQuotaCoordinator(300, 2_000_000)
        self.assertEqual((coordinator.rpm, coordinator.tpm), (240, 1_600_000))

    def test_atomic_rpm_and_concurrency_limits(self):
        now = [0.0]
        coordinator = GeminiQuotaCoordinator(
            5, 1_000_000, headroom=1, max_concurrency=10, clock=lambda: now[0]
        )

        with ThreadPoolExecutor(max_workers=20) as pool:
            reservations = list(pool.map(lambda _: coordinator.try_reserve(1)[0], range(20)))

        granted = [item for item in reservations if item]
        self.assertEqual(len(granted), 5)
        for reservation in granted:
            coordinator.release(reservation)

        self.assertIsNone(coordinator.try_reserve(1)[0])
        now[0] = 60
        self.assertIsNotNone(coordinator.try_reserve(1)[0])

    def test_token_reservations_adapt_to_actual_usage(self):
        coordinator = GeminiQuotaCoordinator(
            100, 200_000, headroom=1, max_concurrency=10
        )
        first, _ = coordinator.try_reserve()
        second, _ = coordinator.try_reserve()
        self.assertIsNone(coordinator.try_reserve()[0])

        coordinator.release(first, actual_tokens=10_000)
        coordinator.release(second, actual_tokens=10_000)
        self.assertEqual(coordinator.token_estimate(), 12_500)
        self.assertIsNotNone(coordinator.try_reserve()[0])

    def test_statement_calls_use_and_release_shared_quota(self):
        from google.genai.errors import ClientError

        class Models:
            def __init__(self):
                self.calls = 0

            def generate_content(self, **kwargs):
                self.calls += 1
                if self.calls == 1:
                    raise ClientError(429, {"message": "test"})
                usage = SimpleNamespace(prompt_token_count=2_000)
                return SimpleNamespace(text="statement", usage_metadata=usage)

        class Client:
            def __init__(self):
                self.models = Models()
                self.closed = False

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.closed = True

        with tempfile.TemporaryDirectory() as workspace:
            pdf = Path(workspace) / "statement.pdf"
            pdf.write_bytes(b"pdf")
            client = Client()
            coordinator = GeminiQuotaCoordinator(
                100, 1_000_000, headroom=1, max_concurrency=1
            )
            with (
                patch("utils.extract.create_genai_client", return_value=client),
                patch("utils.extract.get_gemini_coordinator", return_value=coordinator),
                patch("utils.extract.retry_delay", return_value=0),
                patch("utils.extract.time.sleep"),
            ):
                self.assertEqual(extract_statement(pdf), "statement")

            self.assertEqual(client.models.calls, 2)
            self.assertTrue(client.closed)
            self.assertEqual(coordinator._active, 0)


class ReconciliationTests(unittest.TestCase):
    def test_uses_strict_structured_output_and_returns_parsed_data(self):
        result = {
            "vendor_name": "Vendor",
            "total_amount": "10.00",
            "tax_total": "1.30",
            "date": "07/01/2026",
            "category": "Office Expenses",
            "reconciled": True,
        }
        create = Mock(
            return_value=SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(result)))]
            )
        )
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))

        with patch("utils.helpers.get_openai_client", return_value=client):
            actual = reconcile_with_statement("{}", "statement")

        request = create.call_args.kwargs
        self.assertEqual(actual, result)
        self.assertEqual(request["messages"][0]["role"], "developer")
        self.assertIn("append to the input json", request["messages"][0]["content"])
        self.assertIn("Invoice JSON:\n{}", request["messages"][1]["content"])
        self.assertTrue(request["response_format"]["json_schema"]["strict"])


if __name__ == "__main__":
    unittest.main()
