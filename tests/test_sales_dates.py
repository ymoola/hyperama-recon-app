import tempfile
import unittest
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook

from utils.prompts import build_sales_extraction_prompt
from utils.sales_dates import format_long_dates, normalize_monthly_dates


class SalesDateTests(unittest.TestCase):
    def test_corrects_day_month_swaps_in_monthly_batch(self):
        dates = pd.Series([
            "2026-02-07",
            "2026-07-03",
            "2026-07-04",
            "2026-07-05",
            "2026-09-07",
        ])

        normalized = normalize_monthly_dates(dates)

        self.assertEqual(normalized.dt.strftime("%Y-%m-%d").tolist(), [
            "2026-07-02",
            "2026-07-03",
            "2026-07-04",
            "2026-07-05",
            "2026-07-09",
        ])

    def test_uses_previous_month_when_batch_has_no_dominant_month(self):
        dates = pd.Series(["2026-01-08", "2026-02-08", "2026-03-08"])

        normalized = normalize_monthly_dates(
            dates, reference_date=date(2026, 9, 9)
        )

        self.assertEqual(normalized.dt.strftime("%Y-%m-%d").tolist(), [
            "2026-08-01",
            "2026-08-02",
            "2026-08-03",
        ])

    def test_preserves_clearly_dominant_non_previous_month(self):
        dates = pd.Series(["2026-01-08", "2026-01-09", "2026-01-10"])

        normalized = normalize_monthly_dates(
            dates, reference_date=date(2026, 9, 9)
        )

        self.assertEqual(normalized.dt.strftime("%Y-%m-%d").tolist(), [
            "2026-01-08",
            "2026-01-09",
            "2026-01-10",
        ])

    def test_prompt_includes_dynamic_date_disambiguation_context(self):
        prompt = build_sales_extraction_prompt(date(2026, 9, 9))

        self.assertIn("Today's date is 2026-09-09", prompt)
        self.assertIn("previous month, which is August 2026", prompt)
        self.assertIn("DD/MM/YYYY or MM/DD/YYYY", prompt)
        self.assertIn("1/8/2026 is likely 1 August 2026", prompt)

    def test_excel_uses_real_long_date_cells(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Monthly Report"
        worksheet.append([""])
        worksheet.append(["Date"])
        worksheet.append([datetime(2026, 7, 13)])
        format_long_dates(worksheet)

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as output:
            workbook.save(output.name)
            reopened = load_workbook(output.name)
            cell = reopened["Monthly Report"]["A3"]

        self.assertIsInstance(cell.value, datetime)
        self.assertEqual(cell.value.date().isoformat(), "2026-07-13")
        self.assertEqual(cell.number_format, "dddd, mmmm d, yyyy")


if __name__ == "__main__":
    unittest.main()
